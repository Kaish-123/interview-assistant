"""Robust text extraction from documents (TXT/code/PDF/DOCX + optional others)."""

from __future__ import annotations

import os
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from interview_copilot.shared.logging import get_logger

logger = get_logger("context.extract")

# Plain text / source code
TEXT_EXTS = {
    ".txt",
    ".py",
    ".js",
    ".ts",
    ".jsx",
    ".tsx",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".sh",
    ".bat",
    ".ps1",
    ".sql",
    ".r",
    ".lua",
    ".pl",
    ".yaml",
    ".yml",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".css",
    ".scss",
    ".sass",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".tsv",
    ".ini",
    ".cfg",
    ".toml",
    ".env",
    ".gitignore",
    ".makefile",
    ".log",
    ".conf",
    ".properties",
}

# Skip binaries / media early
SKIP_EXTS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".bmp",
    ".svg",
    ".ico",
    ".webp",
    ".mp3",
    ".mp4",
    ".wav",
    ".avi",
    ".mov",
    ".mkv",
    ".zip",
    ".tar",
    ".gz",
    ".7z",
    ".rar",
    ".exe",
    ".dll",
    ".so",
    ".dylib",
    ".bin",
    ".o",
    ".class",
    ".pyc",
    ".pyo",
    ".ds_store",
    ".db",
    ".sqlite",
}


class ExtractError(RuntimeError):
    """Raised when a file cannot be read as text."""


def _read_plain_text(path: Path) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
            text = "\n".join(pages).strip()
            if text:
                return text
    except Exception as e:
        logger.debug(f"pdfplumber failed for {path.name}: {e}")

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join(p.extract_text() or "" for p in reader.pages).strip()
        if text:
            return text
    except Exception as e:
        logger.debug(f"pypdf failed for {path.name}: {e}")

    raise ExtractError(
        f"PDF read failed for {path.name}. Install pypdf (or pdfplumber) and retry."
    )


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document

        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except ImportError:
        pass
    except Exception as e:
        logger.debug(f"python-docx failed for {path.name}: {e}")

    # Stdlib OOXML fallback
    try:
        wns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path, "r") as z:
            with z.open("word/document.xml") as fxml:
                root = ET.parse(fxml).getroot()
        paras = []
        for para in root.iter(f"{wns}p"):
            runs = [t.text for t in para.iter(f"{wns}t") if t.text]
            paras.append("".join(runs))
        return "\n".join(paras)
    except Exception as e:
        raise ExtractError(f"DOCX read failed for {path.name}: {e}") from e


def _extract_xlsx(path: Path) -> str:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in wb.worksheets:
            rows.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                rows.append("\t".join("" if c is None else str(c) for c in row))
        return "\n".join(rows)
    except Exception as e:
        raise ExtractError(f"Excel read failed for {path.name}: {e}") from e


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        lines: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            lines.append(f"[Slide {i}]")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text)
        return "\n".join(lines)
    except Exception as e:
        raise ExtractError(f"PPTX read failed for {path.name}: {e}") from e


def extract_text_from_file(file_path: str | Path) -> str:
    """
    Extract readable text from a file.

    Supports TXT/code, PDF (pypdf/pdfplumber), DOCX (python-docx or zip/xml),
    XLSX/PPTX when optional libs are installed. Raises ExtractError on failure.
    """
    path = Path(file_path)
    if not path.is_file():
        raise ExtractError(f"File not found: {path}")

    ext = path.suffix.lower()
    basename = path.name

    if ext in SKIP_EXTS or basename == ".DS_Store":
        raise ExtractError(f"Skipped binary/media file: {basename}")

    if ext in TEXT_EXTS or ext == "":
        return _read_plain_text(path)

    if ext == ".pdf":
        return _extract_pdf(path)

    if ext == ".docx":
        return _extract_docx(path)

    if ext in (".xlsx", ".xlsm", ".xls"):
        return _extract_xlsx(path)

    if ext == ".pptx":
        return _extract_pptx(path)

    # Last resort: textract (optional, often unavailable)
    try:
        import textract as _textract

        result = _textract.process(str(path))
        return result.decode("utf-8", errors="replace")
    except Exception as e:
        raise ExtractError(f"Could not extract text from {basename}: {e}") from e
